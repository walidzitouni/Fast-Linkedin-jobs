from __future__ import annotations

import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = Path(__file__).resolve().parent / "jobs.sqlite3"

COLUMNS = [
    "title", "company", "location", "description", "employment_type",
    "date_posted", "job_url", "source", "fetched_at", "search_keyword",
    "first_seen", "last_seen",
]

XLSX_COLUMNS = [
    ("title", "Title", 42),
    ("company", "Company", 26),
    ("location", "Location", 28),
    ("employment_type", "Employment type", 18),
    ("date_posted", "Date posted", 14),
    ("first_seen", "First seen", 20),
    ("last_seen", "Last seen", 20),
    ("search_keyword", "Keyword", 24),
    ("job_url", "Job URL", 46),
    ("description", "Description", 80),
]

DESCRIPTION_CELL_LIMIT = 1000


def _utcnow() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def connect():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS jobs (
            job_url TEXT PRIMARY KEY,
            data_json TEXT NOT NULL,
            first_seen TEXT,
            last_seen TEXT
        )
    """)
    # Migrate databases created before first_seen/last_seen existed.
    cols = {row[1] for row in conn.execute("PRAGMA table_info(jobs)")}
    for col in ("first_seen", "last_seen"):
        if col not in cols:
            conn.execute(f"ALTER TABLE jobs ADD COLUMN {col} TEXT")
    conn.execute(
        "UPDATE jobs SET first_seen = COALESCE(first_seen, ?), "
        "last_seen = COALESCE(last_seen, ?)",
        (_utcnow(), _utcnow()),
    )
    return conn


def upsert_records(records: list[dict]) -> tuple[int, int, list[str]]:
    """Insert new jobs, refresh known ones. Returns (inserted, updated, new_urls)."""
    now = _utcnow()
    inserted = updated = 0
    new_urls: list[str] = []
    with connect() as conn:
        for record in records:
            url = (record.get("job_url") or "").strip()
            if not url:
                continue
            row = conn.execute(
                "SELECT data_json, first_seen FROM jobs WHERE job_url = ?", (url,)
            ).fetchone()
            if row is None:
                record["first_seen"] = now
                record["last_seen"] = now
                conn.execute(
                    "INSERT INTO jobs(job_url, data_json, first_seen, last_seen) "
                    "VALUES (?, ?, ?, ?)",
                    (url, json.dumps(record, ensure_ascii=False), now, now),
                )
                inserted += 1
                new_urls.append(url)
            else:
                existing = json.loads(row[0])
                record["first_seen"] = row[1] or existing.get("first_seen") or now
                record["last_seen"] = now
                # A card refresh has no description/criteria; keep the old ones.
                for field in ("description", "employment_type"):
                    if not record.get(field) and existing.get(field):
                        record[field] = existing[field]
                conn.execute(
                    "UPDATE jobs SET data_json = ?, last_seen = ? WHERE job_url = ?",
                    (json.dumps(record, ensure_ascii=False), now, url),
                )
                updated += 1
    return inserted, updated, new_urls


def save_records(records: list[dict]) -> int:
    """Insert/refresh records; returns how many were new."""
    inserted, _, _ = upsert_records(records)
    return inserted


def load_records() -> list[dict]:
    with connect() as conn:
        rows = conn.execute(
            "SELECT data_json, first_seen, last_seen FROM jobs ORDER BY rowid DESC"
        ).fetchall()
    records = []
    for data_json, first_seen, last_seen in rows:
        record = json.loads(data_json)
        record["first_seen"] = first_seen or record.get("first_seen", "")
        record["last_seen"] = last_seen or record.get("last_seen", "")
        records.append({col: record.get(col, "") for col in COLUMNS})
    return records


def get_new_since(iso_timestamp: str) -> list[dict]:
    with connect() as conn:
        rows = conn.execute(
            "SELECT data_json, first_seen, last_seen FROM jobs "
            "WHERE first_seen >= ? ORDER BY first_seen DESC",
            (iso_timestamp,),
        ).fetchall()
    records = []
    for data_json, first_seen, last_seen in rows:
        record = json.loads(data_json)
        record["first_seen"] = first_seen or ""
        record["last_seen"] = last_seen or ""
        records.append({col: record.get(col, "") for col in COLUMNS})
    return records


def clear_records():
    with connect() as conn:
        conn.execute("DELETE FROM jobs")


def export_csv(records: list[dict]) -> bytes:
    return pd.DataFrame(records, columns=COLUMNS).to_csv(index=False).encode("utf-8-sig")


def export_json(records: list[dict]) -> bytes:
    return json.dumps(records, ensure_ascii=False, indent=2).encode("utf-8")


def _fill_sheet(ws, records: list[dict]) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0A66C2")
    for col_idx, (_, label, width) in enumerate(XLSX_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    url_col = next(i for i, (key, _, _) in enumerate(XLSX_COLUMNS, 1) if key == "job_url")
    desc_col = next(i for i, (key, _, _) in enumerate(XLSX_COLUMNS, 1) if key == "description")

    for row_idx, record in enumerate(records, start=2):
        for col_idx, (key, _, _) in enumerate(XLSX_COLUMNS, start=1):
            value = record.get(key, "") or ""
            if key == "description" and len(value) > DESCRIPTION_CELL_LIMIT:
                value = value[:DESCRIPTION_CELL_LIMIT].rstrip() + "…"
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx == url_col and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            if col_idx == desc_col:
                cell.alignment = Alignment(vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(XLSX_COLUMNS))}{max(1, len(records) + 1)}"


def build_workbook(all_records: list[dict], new_records: list[dict] | None = None) -> Workbook:
    wb = Workbook()
    if new_records is None:
        ws = wb.active
        ws.title = "Jobs"
        _fill_sheet(ws, all_records)
    else:
        ws_new = wb.active
        ws_new.title = "New today"
        _fill_sheet(ws_new, new_records)
        ws_all = wb.create_sheet("All jobs")
        _fill_sheet(ws_all, all_records)
    return wb


def export_xlsx_bytes(records: list[dict]) -> bytes:
    from io import BytesIO

    buffer = BytesIO()
    build_workbook(records).save(buffer)
    return buffer.getvalue()


def write_excel(all_records: list[dict], new_records: list[dict], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(all_records, new_records).save(path)
    return path
